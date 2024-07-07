from docx import Document
from docx.shared import Inches
from docx.oxml.shared import OxmlElement, qn

def shade_cells(cells, shade):
    for cell in cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcVAlign = OxmlElement("w:shd")
        tcVAlign.set(qn("w:fill"), shade)
        tcPr.append(tcVAlign)

document = Document()

table = document.add_table(rows=1, cols=2)

shade_cells([table.cell(0, 0)], "#e7e7e7")

table.rows[0].height = Inches(0.7)

hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Qty'
hdr_cells[1].text = 'Id'

document.add_page_break()

document.save('johnpaul_newman_cv.docx')
